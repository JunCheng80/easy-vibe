# lesson-builder/10-engine/render_lesson_plans_docx.py
# ------------------------------------------------------------
# YAML + KMCC Word templates -> weekly DOCX -> (optional) merged DOCX
#
# Key features:
# 1) week01 uses cover template; week02-16 uses body template
# 2) can render a subset: --from-week 5 --to-week 5
# 3) can skip merge: --no-merge
# 4) can merge only: --merge-only (merge existing _weeks_tmp/*.docx)
#
# Inputs:
#   lesson-builder/00-input/courses/<course>.dl.yaml
#   lesson-builder/00-input/lesson_specs/<course>/yaml/weekXX.yaml
# Templates:
#   lesson-builder/20-templates/kmcc_lesson_plan_master.docx
#   lesson-builder/20-templates/kmcc_lesson_plan_body.docx
# Outputs:
#   <out_dir>/_weeks_tmp/weekXX.lesson_plan.docx
#   <out_path> (merged) unless --no-merge or --merge-only
# ------------------------------------------------------------

import argparse
import glob
import json
import os
import traceback
from pathlib import Path

import yaml
from docxtpl import DocxTemplate
from docx import Document
from docx.shared import Emu
from docxcompose.composer import Composer


def read_yaml(path: Path):
    with open(path, "r", encoding="utf-8") as f:
        return yaml.safe_load(f) or {}


def week_num_from_filename(path: str) -> int:
    base = os.path.basename(path)
    # week05.yaml -> 5
    return int(base.replace("week", "").replace(".yaml", ""))


def normalize_process(process_value):
    """
    Ensure lesson.process is list[dict] with stable keys.
    Prevents common docxtpl failures if someone wrote process as str/dict/etc.
    """
    if process_value is None:
        proc_list = []
    elif isinstance(process_value, list):
        proc_list = process_value
    elif isinstance(process_value, dict):
        proc_list = [process_value]
    elif isinstance(process_value, str):
        proc_list = [{"step": process_value}]
    else:
        proc_list = [{"step": str(process_value)}]

    norm = []
    for item in proc_list:
        if isinstance(item, dict):
            norm.append(item)
        else:
            norm.append({"step": str(item)})

    def fill_defaults(d):
        # Use your YAML keys; if template expects different keys, keep YAML consistent & adjust template.
        return {
            "step": d.get("step", "—"),
            "teach_content": d.get("teach_content", "—"),
            "teacher_action": d.get("teacher_action", "—"),
            "student_action": d.get("student_action", "—"),
            "method": d.get("method", "—"),
            "time": d.get("time", "—"),
        }

    return [fill_defaults(x) for x in norm]


def pad_process_to_4(lesson_obj: dict):
    proc = normalize_process(lesson_obj.get("process"))
    proc = proc[:4]
    while len(proc) < 4:
        proc.append({
            "step": "—",
            "teach_content": "—",
            "teacher_action": "—",
            "student_action": "—",
            "method": "—",
            "time": "—",
        })
    lesson_obj["process"] = proc
    return lesson_obj


def fix_tables_full_width(docx_path: Path):
    """
    Post-process: set table cell widths to reduce "thin column" layout.
    This can be heavy; you can skip it by --skip-fix-tables.
    """
    d = Document(str(docx_path))
    sec = d.sections[0]
    usable = sec.page_width - sec.left_margin - sec.right_margin  # EMU

    for t in d.tables:
        try:
            t.autofit = False
        except Exception:
            pass

        cols = len(t.rows[0].cells) if t.rows else 0
        if cols <= 0:
            continue

        if cols == 6:
            ratios = [0.10, 0.28, 0.18, 0.18, 0.13, 0.13]
            widths = [int(usable * r) for r in ratios]
        elif cols == 4:
            widths = [int(usable / 4)] * 4
        else:
            widths = [int(usable / cols)] * cols

        for row in t.rows:
            for i, cell in enumerate(row.cells):
                w = widths[min(i, len(widths) - 1)]
                cell.width = Emu(w)

    d.save(str(docx_path))


def render_one(template_path: Path, ctx: dict, out_path: Path, skip_fix_tables: bool):
    tpl = DocxTemplate(str(template_path))
    tpl.render(ctx)
    tpl.save(str(out_path))
    if not skip_fix_tables:
        fix_tables_full_width(out_path)


def merge_docx(docx_files, out_path: Path):
    if not docx_files:
        raise RuntimeError("No docx files to merge.")
    base = Document(str(docx_files[0]))
    composer = Composer(base)
    for p in docx_files[1:]:
        composer.append(Document(str(p)))
    composer.save(str(out_path))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--course-slug", default="speech_appdev")
    ap.add_argument("--term-slug", default="2026S3")  # kept for display; not used for path
    ap.add_argument("--out", required=True, help="merged docx output path")
    ap.add_argument("--from-week", type=int, default=1)
    ap.add_argument("--to-week", type=int, default=16)
    ap.add_argument("--no-merge", action="store_true", help="render weekly docx only, do not merge")
    ap.add_argument("--merge-only", action="store_true", help="merge existing _weeks_tmp docx only, do not render")
    ap.add_argument("--continue-on-error", action="store_true", help="skip failed weeks and continue")
    ap.add_argument("--dump-errors", action="store_true", help="dump render_errors.json next to out docx")
    ap.add_argument("--debug", action="store_true")
    ap.add_argument("--skip-fix-tables", action="store_true", help="skip fix_tables_full_width to reduce load")
    args = ap.parse_args()

    repo_root = Path(__file__).resolve().parents[2]

    tpl_cover = repo_root / "lesson-builder/20-templates/kmcc_lesson_plan_master.docx"
    tpl_body = repo_root / "lesson-builder/20-templates/kmcc_lesson_plan_body.docx"
    if not tpl_cover.exists():
        raise FileNotFoundError(f"Missing cover template: {tpl_cover}")
    if not tpl_body.exists():
        raise FileNotFoundError(f"Missing body template: {tpl_body}")

    course_yaml = repo_root / f"lesson-builder/00-input/courses/{args.course_slug}.dl.yaml"
    weeks_dir = repo_root / f"lesson-builder/00-input/lesson_specs/{args.course_slug}/yaml"
    if not course_yaml.exists():
        raise FileNotFoundError(f"Missing course yaml: {course_yaml}")
    if not weeks_dir.exists():
        raise FileNotFoundError(f"Missing weeks dir: {weeks_dir}")

    out_docx = Path(args.out).resolve()
    out_root = out_docx.parent
    out_root.mkdir(parents=True, exist_ok=True)

    tmp_dir = out_root / "_weeks_tmp"
    tmp_dir.mkdir(parents=True, exist_ok=True)

    if args.debug:
        print("[DBG] repo_root =", repo_root)
        print("[DBG] course_yaml =", course_yaml)
        print("[DBG] weeks_dir =", weeks_dir)
        print("[DBG] out_docx =", out_docx)
        print("[DBG] tmp_dir =", tmp_dir)
        print("[DBG] range =", args.from_week, "to", args.to_week)
        print("[DBG] no_merge =", args.no_merge, "| merge_only =", args.merge_only, "| skip_fix_tables =", args.skip_fix_tables)

    # Merge-only mode
    if args.merge_only:
        docx_files = sorted(glob.glob(str(tmp_dir / "week*.lesson_plan.docx")))
        if args.debug:
            print("[DBG] merge-only files:", len(docx_files))
        merge_docx([Path(p) for p in docx_files], out_docx)
        print("[OK] merged:", out_docx)
        return

    course_obj = read_yaml(course_yaml).get("course", {})
    # Build week file list, then filter by range
    all_week_files = sorted(glob.glob(str(weeks_dir / "week*.yaml")), key=week_num_from_filename)
    week_files = []
    for wf in all_week_files:
        wn = week_num_from_filename(wf)
        if args.from_week <= wn <= args.to_week:
            week_files.append(wf)

    if args.debug:
        print("[DBG] week_files selected:", len(week_files))
        for wf in week_files:
            print("   -", wf)

    errors = []
    rendered_files = []

    for wf in week_files:
        wn = week_num_from_filename(wf)
        try:
            raw = read_yaml(Path(wf))
            lesson_obj = raw.get("lesson", {})
            lesson_obj = pad_process_to_4(lesson_obj)

            tpl = tpl_cover if wn == 1 else tpl_body
            out_one = tmp_dir / f"week{wn:02d}.lesson_plan.docx"

            if args.debug:
                print(f"[DBG] week{wn:02d} -> tpl={'cover' if wn==1 else 'body'} -> {out_one}")

            ctx = {"course": course_obj, "lesson": lesson_obj}
            render_one(tpl, ctx, out_one, skip_fix_tables=args.skip_fix_tables)
            rendered_files.append(out_one)
            print("[OK] rendered:", out_one, "(cover)" if wn == 1 else "(body)")

        except Exception as e:
            tb = traceback.format_exc()
            print(f"[ERR] week{wn:02d} failed: {wf}")
            print(tb)
            errors.append({"week": wn, "file": wf, "error": str(e), "traceback": tb})
            if not args.continue_on_error:
                break

    if errors and args.dump_errors:
        err_path = out_root / "render_errors.json"
        err_path.write_text(json.dumps(errors, ensure_ascii=False, indent=2), encoding="utf-8")
        print("[INFO] errors dumped:", err_path)

    # No-merge mode
    if args.no_merge:
        print("[OK] no-merge: weekly docx ready under:", tmp_dir)
        return

    # Merge if we have at least one file
    if not rendered_files:
        raise RuntimeError("No weekly docx rendered successfully; cannot merge.")

    merge_docx(rendered_files, out_docx)
    print("[OK] merged:", out_docx)


if __name__ == "__main__":
    main()