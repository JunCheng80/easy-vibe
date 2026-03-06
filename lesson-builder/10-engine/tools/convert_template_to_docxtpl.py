from docx import Document

# -------------------------
# utils: text normalize
# -------------------------
def norm(s: str) -> str:
    return (s or "").replace("\u3000", " ").replace("\n", " ").strip()

def cell_text(cell) -> str:
    return "\n".join(p.text for p in cell.paragraphs).strip()

# -------------------------
# utils: replace in runs (keep style)
# -------------------------
def replace_in_runs(paragraph, old: str, new: str) -> bool:
    """
    只在 runs 里做替换，不重建段落 => 保住字号/下划线/对齐
    """
    hit = False
    for r in paragraph.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            hit = True
    return hit

def replace_value_after_label_in_paragraph(paragraph, label: str, new_value: str) -> bool:
    """
    适配封面：形如 '专 业：____人工智能大模型____'
    目标：保留 label 及其格式，只把 label 后面的“值”替换为占位符
    做法：把整段文本拼出来，找 label 后面的内容区（通常是冒号后），
    但最终替换还是落在 runs 上（粗暴但稳：用“旧值”定位替换）。
    """
    full = paragraph.text
    if label not in full:
        return False

    # 取 label 后面的原值（尽量保守）
    idx = full.find(label) + len(label)
    old_tail = full[idx:].strip()
    if not old_tail:
        return False

    # 有些封面是 “label + 若干空格 + old_value”
    # 我们用 old_tail 作为替换目标（run 内替换）
    return replace_in_runs(paragraph, old_tail, new_value)

# -------------------------
# utils: table operations
# -------------------------
def replace_pair_label_value_in_row(table, label: str, replacement: str) -> bool:
    """
    适配：| 课程名称 | 人工智能大模型 | 课程编码 | 2EC4A0032 |
    规则：找到某行某格“等于/包含 label”，就把其右侧那个格当作 value 格替换
    额外保护：如果右侧格为空或仍是旧值，直接覆盖；如果右侧格被合并导致引用重复，也尽量覆盖。
    """
    for row in table.rows:
        cells = row.cells
        for i in range(len(cells) - 1):
            left = norm(cell_text(cells[i]))
            right = norm(cell_text(cells[i + 1]))

            # 命中字段名格（通常就是 "课程名称" 或 "课程编码"）
            if left == norm(label) or norm(label) in left:
                # 只改右侧值格，不动左侧字段名格
                cells[i + 1].text = replacement
                return True
    return False

def replace_block_below_heading(table, heading: str, replacement: str) -> bool:
    """
    适配“二、学情分析 / 六、课后作业 / 七、教学反思 / 八、参考资料”
    模式：标题在某一行的某个单元格，正文在下一行同列（或同一大合并单元格）
    => 替换下一行的内容，不动标题行
    """
    for r in range(len(table.rows) - 1):
        row = table.rows[r]
        next_row = table.rows[r + 1]
        for c in range(len(row.cells)):
            if norm(cell_text(row.cells[c])) == norm(heading):
                # 替换下一行同列
                if c < len(next_row.cells):
                    next_row.cells[c].text = replacement
                    return True
                # 兜底：替换下一行第一个 cell
                next_row.cells[0].text = replacement
                return True
    return False

def table_has_process_headers(table) -> bool:
    """
    宽松识别教学过程表：
    - 允许：教师行为/教师活动/教师指导
    - 允许：学生行为/学生活动
    - 允许：时间安排/时间安排（分钟）/时间分配
    """
    if len(table.rows) < 1:
        return False

    header_row = table.rows[0]
    header_txt = " ".join(norm(cell_text(c)) for c in header_row.cells)

    must = ["教学进程", "教学内容"]
    teacher_ok = any(k in header_txt for k in ["教师行为", "教师活动", "教师指导"])
    student_ok = any(k in header_txt for k in ["学生行为", "学生活动"])
    method_ok = any(k in header_txt for k in ["教学方法", "方法"])
    time_ok = any(k in header_txt for k in ["时间安排", "时间分配", "时间"])

    return all(k in header_txt for k in must) and teacher_ok and student_ok and method_ok and time_ok

def convert_process_table(table) -> bool:
    """
    教学过程表（最关键）：
    - 保留表头+1行模板行
    - 模板行注入 docxtpl 循环
    """
    if not table_has_process_headers(table):
        return False
    if len(table.rows) < 2:
        return False

    # 删除第3行及以后（保留表头+1行）
    while len(table.rows) > 2:
        table._tbl.remove(table.rows[2]._tr)

    tpl_row = table.rows[1]

    # 确保至少 6 列
    if len(tpl_row.cells) < 6:
        return False

    tpl_row.cells[0].text = "{% for p in lesson.process %}\n{{ p.step }}"
    tpl_row.cells[1].text = "{{ p.teach_content }}"
    tpl_row.cells[2].text = "{{ p.teacher_action }}"
    tpl_row.cells[3].text = "{{ p.student_action }}"
    tpl_row.cells[4].text = "{{ p.method }}"
    tpl_row.cells[5].text = "{{ p.time }}\n{% endfor %}"
    return True

# -------------------------
# main convert
# -------------------------
def convert(in_docx: str, out_docx: str):
    doc = Document(in_docx)

    # ========== 1) 封面：只替换“值”，不动样式 ==========
    # 你封面里常见行：学 院 / 专 业 / 课 程 名 称 / 授 课 教 师 / 职 称
    cover_map = {
        "学    院：": '{{ course.school | default("—") }}',
        "学 院：": '{{ course.school | default("—") }}',
        "专 业：": '{{ course.major | default("—") }}',
        "课 程 名 称：": '{{ course.name | default("—") }}',
        "授 课 教 师：": '{{ course.teacher | default("—") }}',
        "职    称：": '{{ course.title | default("—") }}',
        "职 称：": '{{ course.title | default("—") }}',
    }

    # 学年/学期：你说固定“第二学期”，我们也给占位符
    # 如果你希望始终写死“第二学期”，把下面替换为 "第二学期"
    term_placeholder = '{{ course.term | default("第二学期") }}'

    for p in doc.paragraphs:
        for label, placeholder in cover_map.items():
            if label in p.text:
                replace_value_after_label_in_paragraph(p, label, placeholder)

        # 处理“第__学期”：尽量不破坏格式，用 run 替换旧字样
        if "第" in p.text and "学期" in p.text:
            # 常见：第一学期/第二学期
            replace_in_runs(p, "第一学期", term_placeholder)
            replace_in_runs(p, "第二学期", term_placeholder)

    # 封面字段有时在表格里（不是段落），也补一轮
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for label, placeholder in cover_map.items():
                        if label in p.text:
                            replace_value_after_label_in_paragraph(p, label, placeholder)
                    if "第一学期" in p.text or "第二学期" in p.text:
                        replace_in_runs(p, "第一学期", term_placeholder)
                        replace_in_runs(p, "第二学期", term_placeholder)

    # ========== 2) 正文：课程信息表（只改值，不改标题） ==========
    for t in doc.tables:
        # 优先：标准“左标题-右值”
        replace_pair_label_value_in_row(t, "课程名称", "{{ course.name }}")
        replace_pair_label_value_in_row(t, "课程编码", '{{ course.code | default("—") }}')
        replace_pair_label_value_in_row(t, "课程编号", '{{ course.code | default("—") }}')
        replace_pair_label_value_in_row(t, "课程性质", '{{ course.type | default("—") }}')
        replace_pair_label_value_in_row(t, "授课对象", '{{ course.audience | default("—") }}')
        replace_pair_label_value_in_row(t, "章节名称", "{{ lesson.chapter_title }}")
        replace_pair_label_value_in_row(t, "教学周", "{{ lesson.week_no }}")
        replace_pair_label_value_in_row(t, "章节学时", "{{ lesson.hours }}")

    # ========== 3) 二/六/七/八：标题下一行内容替换 ==========
    for t in doc.tables:
        replace_block_below_heading(t, "二、学情分析", "{{ lesson.student_analysis }}")
        replace_block_below_heading(t, "六、课后作业", "{{ lesson.homework }}")
        replace_block_below_heading(t, "七、教学反思", "{{ lesson.reflection }}")
        replace_block_below_heading(
            t,
            "八、参考资料",
            "{% for r in lesson.references %}\n{{ loop.index }}. {{ r }}\n{% endfor %}"
        )

    # ========== 4) 三、教学内容分析：内容要三行带标签 ==========
    content_block = (
        "教学基本内容：{{ lesson.content_analysis.basic }}\n"
        "教学重点：{{ lesson.content_analysis.key_points }}\n"
        "教学难点：{{ lesson.content_analysis.difficult_points }}"
    )
    for t in doc.tables:
        replace_block_below_heading(t, "三、教学内容分析", content_block)

    # ========== 5) 四、教学目标：内容要四行带标签 ==========
    obj_block = (
        "价值目标：{{ lesson.objectives.value }}\n"
        "知识目标：{{ lesson.objectives.knowledge }}\n"
        "能力目标：{{ lesson.objectives.ability }}\n"
        "素质目标：{{ lesson.objectives.quality }}"
    )
    for t in doc.tables:
        replace_block_below_heading(t, "四、教学目标", obj_block)

    # ========== 6) 五、教学过程表：注入循环 ==========
    for t in doc.tables:
        convert_process_table(t)

    doc.save(out_docx)
    print("✅ converted template saved to:", out_docx)


if __name__ == "__main__":
    import sys
    if len(sys.argv) != 3:
        print("Usage: python convert_template_to_docxtpl.py <in_docx> <out_docx>")
        sys.exit(1)
    convert(sys.argv[1], sys.argv[2])